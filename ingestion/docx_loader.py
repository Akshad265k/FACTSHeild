"""DOCX loader using python-docx."""
from __future__ import annotations

import io
from pathlib import Path


def load_docx(source: str | bytes | Path) -> str:
    """
    Extract plain text from a .docx file.

    Parameters
    ----------
    source
        File path or raw bytes content of the .docx file.

    Returns
    -------
    str
        All paragraph text joined by newlines.
    """
    try:
        from docx import Document  # type: ignore
    except ImportError as exc:
        raise ImportError(
            "python-docx is required for .docx support. "
            "Install it with: pip install python-docx"
        ) from exc

    if isinstance(source, (str, Path)):
        doc = Document(str(source))
    elif isinstance(source, bytes):
        doc = Document(io.BytesIO(source))
    else:
        raise TypeError(f"Unsupported source type: {type(source)}")

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)
